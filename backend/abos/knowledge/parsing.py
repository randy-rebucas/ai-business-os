"""Document validation and parsing (pipeline steps 1, 3–5).

Every format is turned into a flat list of `Block`s (heading / text / table) with page numbers where the format has
pages, so chunking can respect structure. Format is decided from the bytes, never from the client's content type;
the file extension must agree with it.
"""

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser

from abos.knowledge import ocr

FORMATS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".csv": "csv", ".txt": "txt", ".md": "md",
           ".markdown": "md"}
MIME = {"pdf": "application/pdf", "csv": "text/csv", "txt": "text/plain", "md": "text/markdown", "html": "text/html",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
MAX_ROWS_PER_SHEET = 5000
MAX_PAGES = 1000


class ParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class Block:
    kind: str  # heading | text | table
    text: str
    level: int = 0  # heading level (1 = top)
    page: int | None = None
    rows: list[list[str]] = field(default_factory=list)  # table rows, first row = header


@dataclass
class Parsed:
    blocks: list[Block]
    page_count: int | None = None
    ocr_used: bool = False
    warnings: list[dict] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(b.text) for b in self.blocks)


# ---- validation -----------------------------------------------------------------------------------------------

def _decode_text(data: bytes) -> str:
    if b"\x00" in data[:8192]:
        raise ParseError("unsupported_type", "This looks like a binary file, not text")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("bad_encoding", "Text could not be decoded (use UTF-8)")


def sniff(data: bytes, filename: str) -> str:
    """Return the detected format; raise when the bytes don't match an allowed type or the extension."""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    claimed = FORMATS.get(ext)
    if claimed is None:
        raise ParseError("unsupported_type", "Supported files: PDF, DOCX, XLSX, CSV, TXT and Markdown")
    if not data:
        raise ParseError("empty_file", "The file is empty")
    if data.startswith(b"%PDF-"):
        detected = "pdf"
    elif data.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile as exc:
            raise ParseError("corrupt_file", "The file is damaged") from exc
        detected = "docx" if "word/document.xml" in names else "xlsx" if "xl/workbook.xml" in names else "zip"
    else:
        _decode_text(data)
        detected = claimed if claimed in ("csv", "txt", "md") else "text"
    if detected != claimed:
        raise ParseError("type_mismatch", f"The file content doesn't match its {ext} extension")
    return detected


# ---- normalisation --------------------------------------------------------------------------------------------

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SPACES = re.compile(r"[ \t ]+")


def clean(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = _CTRL.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "\n".join(_SPACES.sub(" ", line).strip() for line in text.split("\n")).strip()


_NUMBERED = re.compile(r"^(\d+(\.\d+)*\.?|[A-Z]\.)\s+\S")


def looks_like_heading(line: str) -> int:
    """0 = not a heading; otherwise a guessed level. Conservative: short, no sentence punctuation."""
    s = line.strip()
    if not s or len(s) > 80 or s.endswith((".", ",", ";", ":")) or len(s.split()) > 12:
        return 0
    m = _NUMBERED.match(s)
    if m:
        return min(1 + m.group(1).rstrip(".").count("."), 3)
    letters = [c for c in s if c.isalpha()]
    if len(letters) >= 3 and all(c.isupper() for c in letters):
        return 1
    words = [w for w in s.split() if w[0].isalpha()]
    if len(words) >= 2 and sum(w[0].isupper() for w in words) / len(words) >= 0.8:
        return 2
    return 0


def _text_blocks(text: str, page: int | None = None) -> list[Block]:
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n", clean(text)):
        lines = [ln for ln in para.split("\n") if ln.strip()]
        if not lines:
            continue
        if len(lines) > 1 and (lvl := looks_like_heading(lines[0])):
            blocks.append(Block("heading", lines[0], lvl, page))
            lines = lines[1:]
        if len(lines) == 1 and (lvl := looks_like_heading(lines[0])):
            blocks.append(Block("heading", lines[0], lvl, page))
            continue
        blocks.append(Block("text", " ".join(lines), 0, page))
    return blocks


def _table_block(rows: list[list[str]], page: int | None = None, title: str | None = None) -> Block | None:
    rows = [[clean(str(c)) if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    text = "\n".join("| " + " | ".join(r) + " |" for r in rows)
    return Block("table", f"{title}\n{text}" if title else text, 0, page, rows)


# ---- format parsers -------------------------------------------------------------------------------------------

def _pdf(data: bytes) -> Parsed:
    from pypdf import PdfReader  # noqa: PLC0415
    from pypdf.errors import PdfReadError  # noqa: PLC0415

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParseError("encrypted", "The PDF is password-protected")
        pages = reader.pages
        if len(pages) > MAX_PAGES:
            raise ParseError("too_many_pages", f"PDFs are limited to {MAX_PAGES} pages")
        texts = [(p.extract_text() or "") for p in pages]
    except PdfReadError as exc:
        raise ParseError("corrupt_file", "The PDF could not be read") from exc
    out = Parsed([], page_count=len(texts))
    empty = [i for i, t in enumerate(texts, 1) if len(t.strip()) < 20]
    if empty:  # step 4: scanned pages need OCR
        if ocr.available():
            for i in empty:
                texts[i - 1] = ocr.pdf_page_text(data, i - 1)
            out.ocr_used = True
        else:
            out.warnings.append({"code": "ocr_unavailable", "message": (
                f"{len(empty)} page(s) have no text layer (scanned?) and OCR is not installed: "
                f"pages {', '.join(map(str, empty[:20]))}{'…' if len(empty) > 20 else ''}")})
    for i, t in enumerate(texts, 1):
        out.blocks.extend(_text_blocks(t, i))
    return out


def _docx(data: bytes) -> Parsed:
    import docx  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415

    document = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            if b := _table_block([[cell.text for cell in row.cells] for row in item.rows]):
                blocks.append(b)
            continue
        text = clean(item.text)
        if not text:
            continue
        style = (item.style.name if item.style is not None else "") or ""
        if style == "Title":  # level 0 so "Heading 1" sections nest under the document title
            blocks.append(Block("heading", text, 0))
        elif style.startswith("Heading") and style[-1:].isdigit():
            blocks.append(Block("heading", text, int(style[-1])))
        else:
            blocks.append(Block("text", text))
    return Parsed(blocks)


def _xlsx(data: bytes) -> Parsed:
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = Parsed([])
    for ws in wb.worksheets:
        rows = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_ROWS_PER_SHEET:
                out.warnings.append({"code": "rows_truncated",
                                     "message": f"Sheet '{ws.title}' has more than {MAX_ROWS_PER_SHEET} rows; "
                                                "only the first were indexed"})
                break
            rows.append(["" if v is None else str(v) for v in row])
        if b := _table_block(rows):
            out.blocks.append(Block("heading", ws.title, 1))
            out.blocks.append(b)
    wb.close()
    return out


def _csv(data: bytes) -> Parsed:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    out = Parsed([])
    if len(rows) > MAX_ROWS_PER_SHEET:
        out.warnings.append({"code": "rows_truncated",
                             "message": f"Only the first {MAX_ROWS_PER_SHEET} rows were indexed"})
        rows = rows[:MAX_ROWS_PER_SHEET]
    if b := _table_block(rows):
        out.blocks.append(b)
    return out


def _markdown(data: bytes) -> Parsed:
    text = clean(_decode_text(data))
    blocks: list[Block] = []
    buf: list[str] = []
    table: list[list[str]] = []

    def flush() -> None:
        nonlocal table
        if buf:
            blocks.extend(_text_blocks("\n".join(buf)))
            buf.clear()
        if table and (b := _table_block([r for r in table if not all(set(c) <= set("-: ") for c in r)])):
            blocks.append(b)
        table = []

    for line in text.split("\n"):
        if m := re.match(r"^(#{1,6})\s+(.*)$", line):
            flush()
            blocks.append(Block("heading", m.group(2).strip(), len(m.group(1))))
        elif line.startswith("|") and line.endswith("|"):
            if buf:
                blocks.extend(_text_blocks("\n".join(buf)))
                buf.clear()
            table.append([c.strip() for c in line.strip("|").split("|")])
        else:
            if table:
                flush()
            buf.append(line)
    flush()
    return Parsed(blocks)


def _plain(data: bytes) -> Parsed:
    return Parsed(_text_blocks(_decode_text(data)))


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg", "nav", "footer", "form", "iframe"}
    BLOCK = {"p", "div", "li", "section", "article", "br", "tr", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.buf: list[str] = []
        self.skip = 0
        self.heading: int = 0
        self.title = ""
        self.in_title = False

    def _flush(self) -> None:
        text = clean(" ".join(self.buf))
        self.buf.clear()
        if text:
            self.blocks.append(Block("heading", text, self.heading) if self.heading else Block("text", text))

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag == "title":
            self.in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = int(tag[1])
        elif tag in self.BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag == "title":
            self.in_title = False
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = 0
        elif tag in self.BLOCK:
            self._flush()

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip:
            self.buf.append(data)


def parse_html(data: bytes) -> tuple[Parsed, str]:
    p = _HTMLText()
    p.feed(_decode_text(data))
    p.close()
    p._flush()
    return Parsed(p.blocks), clean(p.title)


PARSERS = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "csv": _csv, "md": _markdown, "txt": _plain,
           "html": lambda d: parse_html(d)[0]}


def parse(data: bytes, fmt: str) -> Parsed:
    try:
        parsed = PARSERS[fmt](data)
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001 — any library failure on hostile/odd input is a parse failure
        raise ParseError("parse_failed", f"The {fmt.upper()} file could not be parsed ({type(exc).__name__})") from exc
    if not any(b.text.strip() for b in parsed.blocks):
        code = "no_text_ocr_unavailable" if any(w["code"] == "ocr_unavailable" for w in parsed.warnings) else "no_text"
        raise ParseError(code, "No readable text was found in this document" +
                         (" (it appears to be scanned and OCR is not installed)" if code != "no_text" else ""))
    return parsed
